"""Export the manual engine scenarios to an Excel workbook for checking by hand.

Every number in the workbook is an Excel formula built from the blue input cells (except
the hand-calculated and Python-engine values, which are pasted in for comparison), so the
rent-vs-buy formulas are re-computed independently by Excel. After writing, the workbook
is recalculated in Microsoft Excel (if installed) and its values are compared with the
engine; the comparison is added to artefacts/metrics/engine_tests.json.

Run:  python -m src.engine.export_excel
"""
import json
import subprocess
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.config import ARTEFACTS_DIR, METRICS_DIR
from src.engine.manual_scenarios import SCENARIOS
from src.engine.rent_vs_buy import Assumptions, analyse, balance

OUT = ARTEFACTS_DIR / "engine_test_scenarios.xlsx"
FIRST_ROW, YEARS = 21, 30
LAST_ROW = FIRST_ROW + YEARS - 1

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
GREEN = Font(name=FONT, color="008000")
TITLE = Font(name=FONT, bold=True, size=13)
NOTE = Font(name=FONT, italic=True, color="52514E", size=9)
HEAD_FILL = PatternFill("solid", fgColor="E4E3DF")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="B0AFA9"))
MONEY = '$#,##0.00;($#,##0.00);"-"'
PCT = "0.00%"

INPUTS = [  # (row, label, key, number format)
    (4, "Price P ($)", "P", MONEY), (5, "Monthly rent R ($)", "R", MONEY),
    (6, "Down payment d", "d", PCT), (7, "Mortgage rate r", "r", "0.0000%"),
    (8, "Mortgage term T (years)", "T", "0"), (9, "Holding period H (years)", "H", "0"),
    (10, "Property tax tau", "tau", PCT), (11, "Maintenance m", "m", PCT),
    (12, "Insurance h", "h", PCT), (13, "Buying costs cb", "cb", PCT),
    (14, "Selling costs cs", "cs", PCT), (15, "Discount rate k", "k", PCT),
    (16, "Appreciation g", "g", PCT), (17, "Rent growth q", "q", PCT),
]
L_, I_, N_, M_ = "$E$4", "$E$5", "$E$6", "$E$7"
P_, R_, d_, r_, T_, H_ = "$B$4", "$B$5", "$B$6", "$B$7", "$B$8", "$B$9"
tau_, m_, h_, cb_, cs_, k_, g_, q_ = "$B$10", "$B$11", "$B$12", "$B$13", "$B$14", "$B$15", "$B$16", "$B$17"
rng = lambda col: f"{col}{FIRST_ROW}:{col}{LAST_ROW}"

# Result rows: (row, label, key in engine output, formula, number format)
RESULTS = [
    (7, "Monthly payment M", "monthly_payment",
     f"=IF({L_}=0,0,IF({I_}=0,{L_}/{N_},{L_}*{I_}*(1+{I_})^{N_}/((1+{I_})^{N_}-1)))", MONEY),
    (9, "C_B(H) cost of buying (PV)", "cost_buy", f"=INDEX({rng('J')},{H_})", MONEY),
    (10, "C_R(H) cost of renting (PV)", "cost_rent", f"=INDEX({rng('I')},{H_})", MONEY),
    (11, "Delta(H) = C_R - C_B", "delta", f"=INDEX({rng('K')},{H_})", MONEY),
    (12, "Recommendation at H", "recommendation", '=IF(E11>=0,"buy","rent")', "@"),
    (13, "Break-even year (search 1-30)", "break_even_year",
     f'=IF(COUNT({rng("L")})=0,"never",MIN({rng("L")}))', "0"),
    (14, "User cost UC = (k+tau+m+h-g)P", "user_cost", f"=({k_}+{tau_}+{m_}+{h_}-{g_})*{P_}", MONEY),
    (15, "Annual rent 12R", "annual_rent", f"=12*{R_}", MONEY),
    (16, "Loan balance B_H", "balance_H", f"=INDEX({rng('C')},{H_})", MONEY),
]

TABLE = [  # (column, header, formula for row x with year in A{x}, number format)
    ("A", "Year t", None, "0"),
    ("B", "Home value V_t", f"={P_}*(1+{g_})^A{{x}}", MONEY),
    ("C", "Loan balance B_t", f"=IF(OR({L_}=0,A{{x}}>={T_}),0,IF({I_}=0,{L_}-12*A{{x}}*{M_},"
                              f"{L_}*(1+{I_})^(12*A{{x}})-{M_}*((1+{I_})^(12*A{{x}})-1)/{I_}))", MONEY),
    ("D", "Owner cost O_t", f"=IF(A{{x}}<={T_},12*{M_},0)+({tau_}+{m_}+{h_})*{P_}*(1+{g_})^(A{{x}}-1)", MONEY),
    ("E", "PV of O_t", f"=D{{x}}/(1+{k_})^A{{x}}", MONEY),
    ("F", "Sum PV O (1..t)", f"=SUM($E${FIRST_ROW}:E{{x}})", MONEY),
    ("G", "Rent_t", f"=12*{R_}*(1+{q_})^(A{{x}}-1)", MONEY),
    ("H", "PV of Rent_t", f"=G{{x}}/(1+{k_})^A{{x}}", MONEY),
    ("I", "C_R(t)", f"=SUM($H${FIRST_ROW}:H{{x}})", MONEY),
    ("J", "C_B(t)", f"=({d_}+{cb_})*{P_}+F{{x}}-(B{{x}}*(1-{cs_})-C{{x}})/(1+{k_})^A{{x}}", MONEY),
    ("K", "Delta(t) = C_R - C_B", "=I{x}-J{x}", MONEY),
    ("L", "Break-even helper", '=IF(K{x}>=0,A{x},"")', "0"),
]


def engine_output(s: dict) -> dict:
    a = Assumptions(**s["a"])
    r = analyse(s["P"], s["R"], a)
    r["balance_H"] = balance(s["P"], a, a.H)
    r["break_even_year"] = r["break_even_year"] if r["break_even_year"] is not None else "never"
    return r


def scenario_sheet(wb: Workbook, s: dict):
    ws = wb.create_sheet(s["id"])
    values = {"P": s["P"], "R": s["R"], **s["a"]}
    eng = engine_output(s)
    hand = {k: ("never" if v is None else v) for k, v in s["expected"].items()}

    ws["A1"] = f"{s['id']}: {s['name']}"
    ws["A1"].font = TITLE
    ws["A2"] = "Hand calculation: " + s["hand"]
    ws["A2"].font = NOTE
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 48

    ws["A3"], ws["D3"] = "Inputs (blue: edit to explore)", "Results (Excel formulas)"
    for row, label, key, fmt in INPUTS:
        ws.cell(row, 1, label).font = BLACK
        c = ws.cell(row, 2, values[key])
        c.font, c.number_format = BLUE, fmt
    for row, label, f in ((4, "Loan L = (1-d)P", f"=(1-{d_})*{P_}"), (5, "Monthly rate i = r/12", f"={r_}/12"),
                          (6, "Number of payments n = 12T", f"=12*{T_}")):
        ws.cell(row, 4, label).font = BLACK
        c = ws.cell(row, 5, f)
        c.font, c.number_format = BLACK, (MONEY if row == 4 else "0.000000%" if row == 5 else "0")

    for col, head in ((7, "Hand calculation"), (8, "Python engine"), (9, "Excel vs engine"),
                      (10, "Excel vs hand")):
        c = ws.cell(8, col, head)
        c.font, c.fill = BOLD, HEAD_FILL
    for row, label, key, formula, fmt in RESULTS:
        ws.cell(row, 4, label).font = BOLD if key == "delta" else BLACK
        c = ws.cell(row, 5, formula)
        c.font, c.number_format = BLACK, fmt
        if key == "delta":
            c.fill = KEY_FILL
        e = eng[key]
        ce = ws.cell(row, 8, e)
        ce.font, ce.number_format = BLUE, fmt
        text = isinstance(e, str)
        ws.cell(row, 9, f'=IF(E{row}=H{row},"PASS","FAIL")' if text
                else f'=IF(ABS(E{row}-H{row})<0.01,"PASS","FAIL")').font = BLACK
        if key in hand:
            ch = ws.cell(row, 7, hand[key])
            ch.font, ch.number_format = BLUE, fmt
            ws.cell(row, 10, f'=IF(E{row}=G{row},"PASS","FAIL")' if isinstance(hand[key], str)
                    else f'=IF(ABS(E{row}-G{row})<0.01,"PASS","FAIL")').font = BLACK
    ws["G17"] = ("Hand and Python values are pasted for the inputs above; if you edit an input, "
                 "only the Excel column updates.")
    ws["G17"].font = NOTE

    ws.cell(FIRST_ROW - 2, 1, "Year-by-year (selling at the end of year t)").font = BOLD
    for col, head, _, _ in TABLE:
        c = ws[f"{col}{FIRST_ROW - 1}"]
        c.value, c.font, c.fill, c.border = head, BOLD, HEAD_FILL, THIN
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    for t in range(1, YEARS + 1):
        x = FIRST_ROW + t - 1
        for col, _, formula, fmt in TABLE:
            c = ws[f"{col}{x}"]
            c.value = t if formula is None else formula.format(x=x)
            c.font, c.number_format = BLACK, fmt
    widths = {"A": 27, "B": 15, "C": 15, "D": 30, "E": 17, "F": 15, "G": 17, "H": 15,
              "I": 15, "J": 15, "K": 17, "L": 11}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[FIRST_ROW - 1].height = 30
    ws.freeze_panes = f"A{FIRST_ROW}"


def summary_sheet(wb: Workbook):
    ws = wb.create_sheet("Summary", 1)
    heads = ["ID", "Scenario", "Price P", "Rent R", "H", "Payment M", "C_B(H)", "C_R(H)",
             "Delta(H)", "Recommendation", "Break-even year", "User cost", "Checks"]
    for j, head in enumerate(heads, 1):
        c = ws.cell(1, j, head)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, THIN
    for i, s in enumerate(SCENARIOS, 2):
        sid = s["id"]
        ws.cell(i, 1, sid).font = BLACK
        ws.cell(i, 2, s["name"]).font = BLACK
        for j, ref, fmt in ((3, "B4", MONEY), (4, "B5", MONEY), (5, "B9", "0"), (6, "E7", MONEY),
                            (7, "E9", MONEY), (8, "E10", MONEY), (9, "E11", MONEY), (10, "E12", "@"),
                            (11, "E13", "0"), (12, "E14", MONEY)):
            c = ws.cell(i, j, f"={sid}!{ref}")
            c.font, c.number_format = GREEN, fmt
        ws.cell(i, 13, f'=IF(COUNTIF({sid}!I7:J16,"FAIL")=0,"PASS","FAIL")').font = BOLD
    n = len(SCENARIOS) + 1
    ws.cell(n + 2, 2, "All scenarios").font = BOLD
    ws.cell(n + 2, 13, f'=IF(COUNTIF(M2:M{n},"FAIL")=0,"ALL PASS","SOME FAIL")').font = BOLD
    for j, w in enumerate([6, 58, 14, 11, 5, 13, 14, 14, 14, 15, 15, 13, 11], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "C2"


def readme_sheet(wb: Workbook):
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Rent-vs-buy engine: manual test scenarios", TITLE),
        ("Each scenario sheet re-computes the CLAUDE.md formulas with Excel formulas, year by year, "
         "and compares the result with the hand calculation and the Python engine.", BLACK),
        ("", BLACK),
        ("Colour legend", BOLD),
        ("Blue text: hardcoded inputs, hand-calculated values and pasted Python engine values.", BLUE),
        ("Black text: Excel formulas.   Green text: links to another sheet.", BLACK),
        ("Yellow fill: the key result, Delta(H) = C_R(H) - C_B(H) (positive = buying is cheaper).", BLACK),
        ("", BLACK),
        ("Formulas", BOLD),
        ("L = (1-d)P ; i = r/12 ; n = 12T ; M = L i (1+i)^n / ((1+i)^n - 1)  (M = L/n if r = 0, 0 if d = 1)", BLACK),
        ("V_t = P(1+g)^t ; B_t = L(1+i)^(12t) - M((1+i)^(12t) - 1)/i, and B_t = 0 once t >= T", BLACK),
        ("O_t = 12M [t <= T] + (tau+m+h) V_(t-1) ; Rent_t = 12R(1+q)^(t-1)", BLACK),
        ("C_B(H) = (d+cb)P + sum O_t/(1+k)^t - (V_H(1-cs) - B_H)/(1+k)^H ; C_R(H) = sum Rent_t/(1+k)^t", BLACK),
        ("Break-even = smallest H in 1..30 with Delta(H) >= 0 ; User cost UC = (k+tau+m+h-g)P vs 12R", BLACK),
        ("", BLACK),
        ("How to use", BOLD),
        ("Summary: one row per scenario; 'Checks' is PASS when Excel matches the Python engine and "
         "the hand calculation to the cent.", BLACK),
        ("Scenario sheets: edit the blue inputs to explore; the Excel results and the year table "
         "update, the pasted hand and Python values do not.", BLACK),
        ("", BLACK),
        ("Indicative decision support only – not financial or valuation advice", NOTE),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(i, 1, text)
        c.font = font
    ws.column_dimensions["A"].width = 120


def build(path=OUT):
    wb = Workbook()
    readme_sheet(wb)
    for s in SCENARIOS:
        scenario_sheet(wb, s)
    summary_sheet(wb)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def recalculate_in_excel(path) -> bool:
    """Open the workbook in Microsoft Excel, recalculate everything and save (Windows only)."""
    ps = (f"$x = New-Object -ComObject Excel.Application; $x.Visible = $false; "
          f"$x.DisplayAlerts = $false; $wb = $x.Workbooks.Open('{path}'); "
          f"$x.CalculateFull(); $wb.Save(); $wb.Close(); $x.Quit(); "
          f"[void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($x)")
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True,
                       capture_output=True, timeout=120)
        return True
    except Exception as e:
        print("Excel recalculation not possible:", e)
        return False


def verify(path=OUT) -> dict:
    """Read Excel's computed values back and compare them with the engine."""
    wb = load_workbook(path, data_only=True)
    rows, errors = [], 0
    for s in SCENARIOS:
        ws = wb[s["id"]]
        eng = engine_output(s)
        for row, label, key, _, _ in RESULTS:
            xl, py = ws.cell(row, 5).value, eng[key]
            if isinstance(xl, str) and xl.startswith("#"):
                errors += 1
            ok = (xl == py) if isinstance(py, str) or isinstance(xl, str) else abs(xl - py) < 0.01
            rows.append({"scenario": s["id"], "quantity": key, "excel": xl, "engine": py, "pass": ok})
    summary = wb["Summary"]
    overall = summary.cell(len(SCENARIOS) + 3, 13).value
    return {"workbook": str(path.relative_to(ARTEFACTS_DIR.parent)).replace("\\", "/"),
            "recalculated_with": "Microsoft Excel", "formula_errors": errors,
            "values_compared": len(rows), "values_matching": sum(r["pass"] for r in rows),
            "workbook_summary_cell": overall, "mismatches": [r for r in rows if not r["pass"]]}


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    build()
    print(f"Wrote {OUT}")
    if not recalculate_in_excel(OUT):
        return
    check = verify()
    print(json.dumps({k: v for k, v in check.items() if k != "mismatches"}, indent=2))
    for m in check["mismatches"]:
        print("MISMATCH", m)
    path = METRICS_DIR / "engine_tests.json"
    data = json.loads(path.read_text()) if path.exists() else {}
    data["excel_check"] = check
    path.write_text(json.dumps(data, indent=2, default=str))


if __name__ == "__main__":
    main()
